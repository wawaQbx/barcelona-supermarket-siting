"""Standardise every CSV in data_catalog -> data_clean/.
Rules: UTF-8; snake_case ascii headers; no reserved words; each table keeps its ORIGINAL language (ca / es / en, decided per table) and variant
spellings of shared concepts are merged using THAT language's canonical name; district/barri names overwritten from one authoritative dimension;
census-section codes made composite; invalid placeholders (0 codes, impossible coordinates) -> NULL. Originals are never modified."""
import re, io, unicodedata
from pathlib import Path
import numpy as np, pandas as pd
import bcn_naming as N

CAT, OUT = Path("data_catalog"), Path("data_clean")
SKIP_DIRS = ("unused_downloads", "useless")
SECTION_SRC = {"seccio_censal", "codi_seccio_censal", "sec_cens"}

def read_any(p):
    raw = p.read_bytes(); enc = None
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"): txt, enc = raw.decode("utf-16"), "utf-16"      # only with a BOM; utf-16 "decodes" anything
    else:
        for e in ("utf-8-sig", "cp1252"):
            try: txt, enc = raw.decode(e), e; break
            except Exception: continue
    first = txt.split("\n", 1)[0]
    DELIMS = ",;\t|"
    delim = max(DELIMS, key=first.count) if any(d in first for d in DELIMS) else ","
    return pd.read_csv(io.StringIO(txt), sep=delim, dtype=str, keep_default_na=False, na_values=N.NA_TOKENS, low_memory=False), enc

def to_int(s): return pd.to_numeric(s, errors="coerce").astype("Int64")

# ---- authoritative dimensions: names live in ONE place; every table is overwritten from it by code ----
_auth, _ = read_any(CAT / "02_demand_population_income_housing/housing/BarcelonaCiutat_SeccionsCensals_fixed.csv")
_auth.columns = [N.clean_name(c) for c in _auth.columns]
_auth["codi_districte"] = to_int(_auth["codi_districte"]); _auth["codi_barri"] = to_int(_auth["codi_barri"])
DIM_D = _auth[["codi_districte", "nom_districte"]].drop_duplicates().sort_values("codi_districte").reset_index(drop=True)
DIM_N = _auth[["codi_barri", "nom_barri", "codi_districte"]].drop_duplicates().sort_values("codi_barri").reset_index(drop=True)
assert len(DIM_D) == 10 and len(DIM_N) == 73, (len(DIM_D), len(DIM_N))
assert not DIM_D.nom_districte.str.contains("�|ï¿").any() and not DIM_N.nom_barri.str.contains("�|ï¿").any()
(OUT / "dims").mkdir(parents=True, exist_ok=True)
DIM_D.to_csv(OUT / "dims/dim_district.csv", index=False, encoding="utf-8")
DIM_N.to_csv(OUT / "dims/dim_neighborhood.csv", index=False, encoding="utf-8")
MAP_D = dict(zip(DIM_D.codi_districte, DIM_D.nom_districte)); MAP_N = dict(zip(DIM_N.codi_barri, DIM_N.nom_barri))

def out_path(rel):
    nz = lambda x: re.sub(r"[^a-z0-9_]+", "_", unicodedata.normalize("NFKD", x).encode("ascii", "ignore").decode().lower())
    return OUT.joinpath(*[nz(x) for x in rel.parts[:-1]], re.sub(r"[^a-z0-9_.]+", "_", unicodedata.normalize("NFKD", rel.name).encode("ascii", "ignore").decode().lower()))

_g, _ = read_any(CAT / "03_commercial_census/census_editions/census_2022.csv")
GOOD_EIX = _g["Nom_Eix"].dropna().unique().tolist()      # clean street-axis names, used to repair U+FFFD in other editions
dic, report = [], []
for f in sorted(CAT.rglob("*.csv")):
    if any(d in f.parts for d in SKIP_DIRS): continue
    df, enc = read_any(f); n0 = len(df); notes = []
    if f.name == "census_2014.csv":
        for c in ("X_ED50", "Y_ED50"): df[c] = pd.to_numeric(df[c].astype(str).str.replace(",", ".", regex=False), errors="coerce")
        df = df.drop(columns=["X", "Y"]); notes.append("dropped corrupted X/Y (scientific-notation damage); ED50 decimal comma fixed; coordinates are EPSG:23031")
    # text hygiene: trim, safe mojibake repair
    fixed_moj = 0
    for col in df.columns:
        if pd.api.types.is_object_dtype(df[col]) or pd.api.types.is_string_dtype(df[col]):
            s = df[col].str.strip(); fx = s.map(N.fix_mojibake); fixed_moj += int(((fx != s) & s.notna()).sum()); df[col] = fx
    if fixed_moj: notes.append(f"mojibake repaired (provably correct round-trip): {fixed_moj} values")
    orig_cols = list(df.columns); lang = N.detect_lang(f.stem, orig_cols); K = N.KEYS[lang]
    names = N.std_cols(orig_cols, lang, f.stem); mapping = dict(zip(orig_cols, names)); df.columns = names
    df = N.merge_dup_keys(df, notes)
    drop = [c for c in N.DROP_COLS.get(f.stem.lower(), []) if c in df.columns]
    if drop: df = df.drop(columns=drop); notes.append(f"dropped exact duplicate columns {drop}")
    for col in [c for c in df.columns if c in N.INT_COLS]: df[col] = pd.to_numeric(df[col], errors="coerce").round().astype("Int64")
    for col in [c for c in df.columns if c in N.YYYYMM_COLS]:
        v = pd.to_numeric(df[col], errors="coerce").astype("Int64").astype(str)
        df[col] = v.where(v.str.fullmatch(r"\d{6}"), None).str.replace(r"^(\d{4})(\d{2})$", r"\1-\2-01", regex=True); notes.append(f"{col}: yyyymm -> date (first day of the month)")
    sn = [col for col in df.columns if col.startswith("sn_")]
    for col in sn: df[col] = N.sn_to_bool(df[col], inverted=f.stem.lower() in N.SN_INVERTED_STEMS)
    if sn: notes.append(f"{len(sn)} sn_* flags -> boolean (true = Si)" + (" [INVERTED source coding]" if f.stem.lower() in N.SN_INVERTED_STEMS else ""))
    for col in list(df.columns):
        if pd.api.types.is_object_dtype(df[col]) or pd.api.types.is_string_dtype(df[col]):
            new, ok, nb = N.dmy_to_iso(df[col])
            if ok: df[col] = new; notes.append(f"{col}: dd/mm/yyyy -> ISO date" + (f" ({nb} impossible placeholder dates -> NULL)" if nb else ""))
    if "nom_eix" in df.columns:
        df["nom_eix"], nfix, nbad = N.repair_fffd(df["nom_eix"], GOOD_EIX)
        if nbad: notes.append(f"nom_eix: {nfix} of {nbad} damaged names repaired from clean edition; {nbad - nfix} unrecoverable")
    # codes -> integers; names overwritten from the authoritative dimension (only where both columns exist)
    for code, name, mp in ((K["dcode"], K["dname"], MAP_D), (K["bcode"], K["bname"], MAP_N)):
        if code in df.columns:
            df[code] = to_int(df[code])
            if name in df.columns:
                new = df[code].map(mp); changed = int((new.notna() & (new != df[name])).sum()); df[name] = new.where(new.notna(), df[name])
                if changed: notes.append(f"{name}: {changed} rows overwritten from authoritative dimension")
            bad = (df[code] <= 0).fillna(False)
            if bad.any(): df.loc[bad, code] = pd.NA; notes.append(f"{code}: {int(bad.sum())} zero/negative -> NULL")
    if "lat" in df.columns and "lon" in df.columns:
        la, lo = pd.to_numeric(df["lat"], errors="coerce"), pd.to_numeric(df["lon"], errors="coerce")
        bad = ~((la.between(39, 43)) & (lo.between(0, 4))) & (la.notna() | lo.notna())
        if bad.any(): df.loc[bad, ["lat", "lon"]] = np.nan; notes.append(f"lat/lon: {int(bad.sum())} implausible pairs -> NULL")
    for col in list(df.columns):
        if pd.api.types.is_object_dtype(df[col]) or pd.api.types.is_string_dtype(df[col]):
            new, bad = N.stray_markers_to_na(df[col], col)
            if bad: df[col] = new; notes.append(f"{col}: placeholder words {bad} -> NULL")
    df = N.compose_section(df, notes)
    p = out_path(f.relative_to(CAT)); p.parent.mkdir(parents=True, exist_ok=True); df.to_csv(p, index=False, encoding="utf-8")
    assert len(df) == n0, f"row count changed in {f}"
    rel = str(f.relative_to(CAT)).replace("\\", "/")
    for o, n in mapping.items(): dic.append((rel, lang, o, n))
    report.append((rel, str(p.relative_to(OUT)).replace("\\", "/"), lang, n0, len(orig_cols), df.shape[1], enc, "; ".join(notes)))

# ---------- GTFS (TMB): international standard field names are kept as-is; only compact dates become ISO dates ----------
GTFS_DIR = CAT / "04_transport_traffic/gtfs_tmb"
for f in sorted(GTFS_DIR.glob("*.txt")):
    g = pd.read_csv(f, dtype=str, keep_default_na=False, na_values=[""], encoding="utf-8-sig"); n0 = len(g); notes = ["GTFS standard field names kept"]
    for col in g.columns:
        g[col] = g[col].str.strip()
        if col in ("date", "start_date", "end_date", "feed_start_date", "feed_end_date"):
            new, ok = N.yyyymmdd_to_iso(g[col])
            if ok: g[col] = new; notes.append(f"{col}: yyyymmdd -> ISO date")
    p = OUT / "04_transport_traffic/gtfs_tmb" / (f.stem + ".csv"); p.parent.mkdir(parents=True, exist_ok=True); g.to_csv(p, index=False, encoding="utf-8")
    assert len(g) == n0
    rel = str(f.relative_to(CAT)).replace("\\", "/")
    for col in g.columns: dic.append((rel, "en", col, col))
    report.append((rel, str(p.relative_to(OUT)).replace("\\", "/"), "en", n0, g.shape[1], g.shape[1], "utf-8-sig", "; ".join(notes)))

# ---------- rental contracts per year (xlsx: city 2000-2025, 10 districts 2000-2025, 73 barris 2013-2025) -> long table ----------
import openpyxl
X = CAT / "02_demand_population_income_housing/housing/anual_bcn_contractes.xlsx"
rows = list(openpyxl.load_workbook(X, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
hdr_i = next(i for i, r in enumerate(rows) if r and str(r[0]).strip() == "Codi")
years = [int(str(y).strip()) for y in rows[hdr_i][2:] if y is not None]
level, out = None, []
for r in rows[hdr_i + 1:]:
    if not r or all(v is None for v in r): continue
    label = str(r[1]).strip() if r[1] is not None else ""
    if r[0] is None:
        lo = label.lower()
        if lo.startswith("barcelona"): level = "municipi"
        elif lo.startswith("districte"): level = "districte"; continue
        elif lo.startswith("barri"): level = "barri"; continue
        else: continue                                            # footnotes / source lines
    if r[0] is not None and not str(r[0]).strip().isdigit(): continue      # footnote text in the code column
    code = int(r[0]) if r[0] is not None else None
    for y, v in zip(years, r[2:2 + len(years)]):
        if v is None: continue
        out.append(dict(nivell=level, codi_districte=(code if level == "districte" else (DIM_N.set_index("codi_barri").codi_districte.get(code) if level == "barri" else None)),
                        codi_barri=(code if level == "barri" else None), nom=label, any_=y, nombre_contractes=int(v)))
lc = pd.DataFrame(out)
lc.loc[lc.nivell == "barri", "nom"] = lc.loc[lc.nivell == "barri", "codi_barri"].map(MAP_N)          # authoritative names
lc.loc[lc.nivell == "districte", "nom"] = lc.loc[lc.nivell == "districte", "codi_districte"].map(MAP_D)
for col in ("codi_districte", "codi_barri"): lc[col] = lc[col].astype("Int64")
chk = lc[lc.nivell == "barri"].groupby("any_").nombre_contractes.sum(); city = lc[lc.nivell == "municipi"].set_index("any_").nombre_contractes
p = OUT / "02_demand_population_income_housing/housing/contractes_lloguer_anual.csv"; lc.to_csv(p, index=False, encoding="utf-8")
notes = [f"long format: {lc.nivell.value_counts().to_dict()} rows", f"barris present {lc[lc.nivell=='barri'].codi_barri.nunique()}, years {lc[lc.nivell=='barri'].any_.min()}-{lc[lc.nivell=='barri'].any_.max()}",
         f"sum of barris vs city 2025: {int(chk.get(2025, 0))} / {int(city.get(2025, 0))}"]
for col in lc.columns: dic.append(("02_demand_population_income_housing/housing/anual_bcn_contractes.xlsx", "ca", col, col))
report.append(("02_demand_population_income_housing/housing/anual_bcn_contractes.xlsx", str(p.relative_to(OUT)).replace("\\", "/"), "ca", len(lc), len(years) + 2, lc.shape[1], "xlsx", "; ".join(notes)))

pd.DataFrame(dic, columns=["source_file", "language", "original_column", "clean_column"]).to_csv(OUT / "COLUMN_DICTIONARY.csv", index=False)
pd.DataFrame(report, columns=["source_file", "clean_file", "language", "rows", "cols_in", "cols_out", "encoding", "notes"]).to_csv(OUT / "FILE_REPORT.csv", index=False)
print(len(report), "files cleaned;", pd.Series([r[2] for r in report]).value_counts().to_dict())
